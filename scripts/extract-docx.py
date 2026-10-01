# -*- coding: utf-8 -*-
"""
把 .docx 完整导出为 UTF-8 markdown（含表格），便于阅读与检索。

用法：
    python scripts/extract-docx.py <输入.docx> [输出.md]

不传输出路径时，默认写到 .agent/<原文件名>-原文提取.md
"""
import sys
import os
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


def iter_block_items(parent):
    """按文档顺序遍历段落与表格"""
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    src = sys.argv[1]
    if not os.path.isfile(src):
        print(f"找不到文件: {src}")
        sys.exit(1)

    if len(sys.argv) >= 3:
        out = sys.argv[2]
    else:
        base = os.path.splitext(os.path.basename(src))[0]
        out = os.path.join(".agent", f"{base}-原文提取.md")

    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    doc = Document(src)

    lines = [
        f"# {os.path.basename(src)} —— 原文提取\n",
        f"> 来源：{src}\n",
        f"> 段落 {len(doc.paragraphs)} ｜ 表格 {len(doc.tables)}\n",
    ]

    tbl_no = 0
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            t = block.text.strip()
            if not t:
                continue
            style = block.style.name if block.style else ""
            if "Heading 1" in style or "标题 1" in style:
                lines.append(f"\n## {t}\n")
            elif "Heading 2" in style or "标题 2" in style:
                lines.append(f"\n### {t}\n")
            elif "Heading 3" in style or "标题 3" in style:
                lines.append(f"\n#### {t}\n")
            else:
                lines.append(t)
        else:
            tbl_no += 1
            lines.append(f"\n**[表 {tbl_no}]**\n")
            for row in block.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")

    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    print(f"OK -> {out}")
    print(f"表格数: {tbl_no}")
    print(f"字符数: {sum(len(l) for l in lines)}")


if __name__ == "__main__":
    main()
